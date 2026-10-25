"""Word export (python-docx): title page, Introduction, Aims, Background, Methods ۳-۱ … ۳-۵, References.

For Persian every paragraph and run is right-to-left (w:bidi, w:rtl), the complex-script font is "B Nazanin" (with
"Vazirmatn" as its alternative in the font table), the Latin font is Times New Roman and text is justified.
References stay left-to-right because they are English.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree

from engine.titles import DISCLAIMER, title

FONT_LATIN = "Times New Roman"
FONT_PERSIAN = "B Nazanin"
FONT_PERSIAN_ALT = "Vazirmatn"
BODY_PT = 12

# Elements that must come after <w:bidi> inside <w:pPr> and after <w:rtl> inside <w:rPr> (OOXML order).
AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip
AFTER_RTL = ("w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
AFTER_SZCS = ("w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign") + ("w:rtl",) + AFTER_RTL


def set_paragraph_rtl(paragraph: Paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.insert_element_before(OxmlElement("w:bidi"), *AFTER_BIDI)


def style_run(run: Any, rtl: bool, size: int = BODY_PT, bold: bool = False) -> None:
    """Fonts for both scripts, size for both scripts, and w:rtl for Persian runs."""
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.bold = bold or None
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    fonts.set(qn("w:cs"), FONT_PERSIAN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    if r_pr.find(qn("w:szCs")) is None:
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), str(size * 2))
        r_pr.insert_element_before(size_cs, *AFTER_SZCS)
    if bold and r_pr.find(qn("w:bCs")) is None:
        r_pr.find(qn("w:b")).addnext(OxmlElement("w:bCs"))
    if rtl and r_pr.find(qn("w:rtl")) is None:
        r_pr.insert_element_before(OxmlElement("w:rtl"), *AFTER_RTL)


def add_paragraph(doc: Any, text: str, rtl: bool, size: int = BODY_PT, bold: bool = False,
                  justify: bool = True, style: str | None = None) -> Paragraph:  # fmt: skip
    paragraph = doc.add_paragraph(style=style)
    if rtl:
        set_paragraph_rtl(paragraph)
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_run(paragraph.add_run(text), rtl, size, bold)
    return paragraph


def add_heading(doc: Any, text: str, rtl: bool, level: int = 1) -> Paragraph:
    paragraph = doc.add_heading(level=level)
    if rtl:
        set_paragraph_rtl(paragraph)
    style_run(paragraph.add_run(text), rtl, {0: 20, 1: 16, 2: 13}.get(level, 12), bold=True)
    return paragraph


def add_font_alternative(doc: Any) -> None:
    """Declare B Nazanin in the font table with Vazirmatn as its alternative (w:altName)."""
    for part in doc.part.package.iter_parts():
        if str(part.partname) != "/word/fontTable.xml":
            continue
        root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
        font = etree.SubElement(root, qn("w:font"))
        font.set(qn("w:name"), FONT_PERSIAN)
        alt = etree.SubElement(font, qn("w:altName"))
        alt.set(qn("w:val"), FONT_PERSIAN_ALT)
        charset = etree.SubElement(font, qn("w:charset"))
        charset.set(qn("w:val"), "B2")  # Arabic
        if not hasattr(part, "_element"):
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return


def set_section_rtl(doc: Any) -> None:
    """Right-to-left page layout for Persian documents (w:bidi in the section properties)."""
    sect_pr = doc.sections[0]._sectPr
    if sect_pr.find(qn("w:bidi")) is None:
        sect_pr.insert_element_before(OxmlElement("w:bidi"), "w:rtlGutter", "w:docGrid", "w:printerSettings")


def build_docx(result: dict[str, Any]) -> bytes:
    """Render a finished proposal (the pipeline's result) as a .docx file."""
    lang = result.get("language", "fa")
    rtl = lang == "fa"
    doc = Document()
    if rtl:
        add_font_alternative(doc)
        set_section_rtl(doc)
    write_title_page(doc, result, rtl)
    sections = result.get("sections", {})
    for key, writer in (("introduction", write_text_section), ("aims", write_aims),
                        ("background", write_background), ("methods", write_methods)):  # fmt: skip
        section = sections.get(key)
        if section and section.get("status") == "done":
            writer(doc, section, lang, rtl)
    write_references(doc, result.get("references", []), lang, rtl)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def write_title_page(doc: Any, result: dict[str, Any], rtl: bool) -> None:
    heading = doc.add_heading(level=0)
    if rtl:
        set_paragraph_rtl(heading)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(heading.add_run(result.get("title", "")), rtl, 20, bold=True)
    created = (result.get("generated_at") or date.today().isoformat())[:10]
    add_paragraph(doc, created, rtl=False, justify=False).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(doc, DISCLAIMER, rtl=False, size=10, justify=False).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def write_text_section(doc: Any, section: dict[str, Any], lang: str, rtl: bool) -> None:
    add_heading(doc, section.get("title") or title("introduction", lang), rtl)
    for block in (section.get("text") or "").split("\n\n"):
        if block.strip():
            add_paragraph(doc, block.strip(), rtl)


def write_aims(doc: Any, section: dict[str, Any], lang: str, rtl: bool) -> None:
    add_heading(doc, section.get("title") or title("aims", lang), rtl)
    add_heading(doc, title("main_aim", lang), rtl, level=2)
    add_paragraph(doc, section.get("main_aim", ""), rtl)
    hypotheses = [h["text"] if h.get("applicable", True) else f"{title('not_applicable', lang)}: {h['text']}"
                  for h in section.get("hypotheses", [])]  # fmt: skip
    for key, items in (("secondary_aims", section.get("secondary_aims", [])), ("hypotheses", hypotheses),
                       ("questions", section.get("questions", [])),
                       ("practical_aims", section.get("practical_aims", []))):  # fmt: skip
        if not items:
            continue
        add_heading(doc, title(key, lang), rtl, level=2)
        for number, item in enumerate(items, start=1):
            add_paragraph(doc, f"{number}. {item}", rtl)


def write_background(doc: Any, section: dict[str, Any], lang: str, rtl: bool) -> None:
    add_heading(doc, section.get("title") or title("background", lang), rtl)
    for paragraph in section.get("paragraphs", []):
        add_paragraph(doc, paragraph["text"], rtl)


def write_methods(doc: Any, section: dict[str, Any], lang: str, rtl: bool) -> None:
    add_heading(doc, section.get("title") or title("methods", lang), rtl)
    for sub in section.get("subsections", []):
        if not sub.get("text"):
            continue
        add_heading(doc, sub["title"], rtl, level=2)
        for block in sub["text"].split("\n"):
            if block.strip():
                add_paragraph(doc, block.strip(), rtl)


def write_references(doc: Any, references: list[dict[str, Any]], lang: str, rtl: bool) -> None:
    if not references:
        return
    add_heading(doc, title("references", lang), rtl)
    for reference in references:
        # References are English: left-to-right and left-aligned even in a Persian document.
        paragraph = add_paragraph(doc, f"{reference['number']}. {reference['text']}", rtl=False, justify=False)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
