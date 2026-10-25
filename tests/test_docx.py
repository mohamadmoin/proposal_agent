import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from engine.docx_export import build_docx
from engine.pipeline import RunOptions, run
from engine.schemas import ProposalRequest
from engine.store import MemoryStore
from engine.titles import DISCLAIMER
from tests.test_pipeline import FA_TITLE


@pytest.fixture
def fa_docx(fake_llm, offline_pubmed, now):
    result = run(ProposalRequest(title=FA_TITLE), MemoryStore(), fake_llm, offline_pubmed, now, RunOptions())
    return result, Document(io.BytesIO(build_docx(result)))


def is_rtl(paragraph):
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.find(qn("w:bidi")) is not None


def test_docx_opens_with_title_page_and_sections(fa_docx):
    result, doc = fa_docx
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == FA_TITLE
    assert DISCLAIMER in texts
    for heading in ("مقدمه", "اهداف، فرضیه‌ها و سؤالات پژوهش", "پیشینه پژوهش", "۳-۱ روش اجرا",
                    "۳-۵ روش تجزیه و تحلیل داده‌ها", "منابع"):  # fmt: skip
        assert heading in texts
    order = [texts.index(h) for h in ("مقدمه", "اهداف، فرضیه‌ها و سؤالات پژوهش", "پیشینه پژوهش", "۳-۱ روش اجرا")]
    assert order == sorted(order)


def test_persian_body_is_rtl_with_fonts(fa_docx):
    _, doc = fa_docx
    body = next(p for p in doc.paragraphs if p.text.startswith("در مطالعه"))
    assert is_rtl(body)
    assert body.alignment == 3  # justified
    run = body.runs[0]._r
    assert run.rPr.find(qn("w:rtl")) is not None
    fonts = run.rPr.find(qn("w:rFonts"))
    assert fonts.get(qn("w:cs")) == "B Nazanin" and fonts.get(qn("w:ascii")) == "Times New Roman"


def test_references_are_left_to_right(fa_docx):
    result, doc = fa_docx
    texts = [p.text for p in doc.paragraphs]
    start = texts.index("منابع")
    references = doc.paragraphs[start + 1 :]
    assert len(references) == len(result["references"])
    assert all(not is_rtl(p) for p in references)
    assert all(p.runs[0]._r.rPr.find(qn("w:rtl")) is None for p in references)
    assert references[0].text.startswith("1. ")


def test_font_table_declares_the_alternative(fa_docx):
    _, doc = fa_docx
    part = next(p for p in doc.part.package.iter_parts() if str(p.partname) == "/word/fontTable.xml")
    xml = part.blob.decode("utf-8") if not hasattr(part, "_element") else part._element.xml
    assert 'w:name="B Nazanin"' in xml and 'w:val="Vazirmatn"' in xml


def test_english_docx_is_not_rtl(fake_llm, offline_pubmed, now):
    request = ProposalRequest(title="Prevalence of vitamin D deficiency among pregnant women", language="en")
    result = run(request, MemoryStore(), fake_llm, offline_pubmed, now, RunOptions())
    doc = Document(io.BytesIO(build_docx(result)))
    assert not any(is_rtl(p) for p in doc.paragraphs)
    assert "3.3 Sample size and sampling" in [p.text for p in doc.paragraphs]
